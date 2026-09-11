"""Regression checks for document links and lossless sheet metadata refreshes.

Run: python3 -m unittest discover -s scripts -p 'test_parse_disciplines_gdoc.py'
"""
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

from parse_disciplines_gdoc import (
    cell_text_links, parse_cell_resources, parse_docx, preserve_sheet_columns,
)


def hyperlink(paragraph, label, url):
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


class DocumentConversionTests(unittest.TestCase):
    def test_embedded_label_and_full_description(self):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        p = cell.paragraphs[0]
        p.add_run("A useful guide. ")
        hyperlink(p, "Guide title", "https://example.org/guide")
        p.add_run(" with teaching examples.")
        text, links = cell_text_links(cell)
        self.assertEqual(parse_cell_resources("Open Methods", text, links), [{
            "title": "A useful guide. Guide title with teaching examples",
            "link": "https://example.org/guide", "category": "Open Methods",
        }])

    def test_partial_url_uses_hyperlink_target(self):
        text = "Paper. https://doi.org/10.1234/truncate"
        spans = [(7, len(text), "https://doi.org/10.1234/truncated")]
        self.assertEqual(parse_cell_resources("General", text, spans)[0]["link"], "https://doi.org/10.1234/truncated")

    def test_split_url_ignores_stale_prefix_target(self):
        text = "Paper. https://doi.org/10.1234/paper"
        spans = [(7, 23, "https://doi.org/10.1234/wrong"), (23, len(text), "https://doi.org/10.1234/paper")]
        self.assertEqual(len(parse_cell_resources("General", text, spans)), 1)
        self.assertEqual(parse_cell_resources("General", text, spans)[0]["link"], "https://doi.org/10.1234/paper")

    def test_repeated_links_deduplicate_but_distinct_links_survive(self):
        rows = parse_cell_resources("Preregistration", "Guide. https://example.org/ahttps://example.org/b https://example.org/a")
        self.assertEqual([r["link"] for r in rows], ["https://example.org/a", "https://example.org/b"])
        self.assertTrue(all(r["title"] == "Guide" and r["category"] == "Open Methods" for r in rows))

    def test_bare_doi_and_missing_link(self):
        self.assertEqual(parse_cell_resources("General", "Paper. doi:10.1234/paper")[0], {
            "title": "Paper", "link": "https://doi.org/10.1234/paper", "category": "General",
        })
        self.assertEqual(parse_cell_resources("", "Unlinked resource")[0]["link"], "")

    def test_repair_and_safelink(self):
        rows = parse_cell_resources("General", "Journal. https://eur01.safelinks.protection.outlook.com/?url=https%3A%2F%2Fwww.degruyter.com%2Fview%2Fjournals%2Feng%2Feng-overview.xml&data=tracking")
        self.assertEqual(rows[0]["link"], "https://www.degruyterbrill.com/journal/key/eng/html")

    def test_encoded_trailing_space_is_removed(self):
        rows = parse_cell_resources("General", "Paper. https://doi.org/10.1234/paper%20")
        self.assertEqual(rows[0]["link"], "https://doi.org/10.1234/paper")

    def test_parentheses_in_doi_are_not_left_in_title(self):
        url = "https://doi.org/10.1016/S0313-5926(09)50047-1"
        self.assertEqual(parse_cell_resources("General", "Paper. " + url)[0], {
            "title": "Paper", "link": url, "category": "General",
        })
        self.assertEqual(parse_cell_resources("General", "Paper (https://example.org/paper)")[0], {
            "title": "Paper", "link": "https://example.org/paper", "category": "General",
        })

    def test_document_without_cover_and_example_hyperlinks(self):
        doc = Document()
        doc.add_heading("Natural Sciences", 1)
        doc.add_heading("Chemistry", 2)
        doc.add_heading("Examples of open research practices", 3)
        p = doc.add_paragraph("Open Data: ")
        hyperlink(p, "Example", "https://example.org/data")
        doc.add_heading("Table of resources", 3)
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Category"
        table.cell(0, 1).text = "Resources"
        table.cell(1, 0).text = "General"
        table.cell(1, 1).text = "Useful guide. https://example.org/guide"
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "source.docx"
            doc.save(path)
            discipline = parse_docx(path)["fields"][0]["disciplines"][0]
        self.assertEqual(discipline["examples"], "**Open Data** [Example](https://example.org/data)")
        self.assertEqual(len(discipline["resources"]), 1)

    def test_metadata_follows_reordered_rows(self):
        existing = [["Field", "Discipline", "Examples", "Leads"],
                    ["Science", "A", "old", "Alice"], ["Science", "B", "old", "Bob"]]
        fresh = [["Field", "Discipline", "Examples"],
                 ["Science", "B", "new"], ["Science", "A", "new"], ["Science", "C", "new"]]
        rows = preserve_sheet_columns("Disciplines", fresh, existing)
        self.assertEqual([r[-1] for r in rows[1:]], ["Bob", "Alice", ""])
        fields = preserve_sheet_columns("Fields", [["Name", "Summary", "Show"], ["Science", "", "TRUE"]],
                                        [["Name", "Summary", "Show", "Leads"], ["Science", "Curated", "FALSE", "Alice"]])
        self.assertEqual(fields[1], ["Science", "Curated", "FALSE", "Alice"])

    def test_resource_backup_follows_link_after_title_changes(self):
        old = [["Discipline", "Title", "Link", "Category", "Category (pre-cleanup backup)"],
               ["A", "Old", "https://example.org/", "General", "Original category"]]
        new = [["Discipline", "Title", "Link", "Category"], ["A", "Better title", "https://example.org", "General"]]
        self.assertEqual(preserve_sheet_columns("Resources", new, old)[1][-1], "Original category")


if __name__ == "__main__":
    unittest.main()
